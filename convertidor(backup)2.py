import openpyxl
from pyxlsb import open_workbook as open_xlsb
from pyxlsb import convert_date
import sys
from os import system, name, path
import datetime
import pymysql

# ------ PIP ------
# pip install pyxlsb https://pypi.org/project/pyxlsb/
# pip install openpyxl https://pypi.org/project/openpyxl/

class Convertidor:
    
    def __init__(self):
        #--------BASE DE DATOS---------
        self.BD     = "flaskcrud"
        #------------TABLA--------------
        self.tabla  = "prueba2"

    def extraerColumnasFechas(self,Row):
        columnas_fecha = []
        columna = 0
        for r in Row:
            if ( 'FECHA' in str(r).encode(encoding="utf-8",errors="xmlcharrefreplace").decode("utf-8").upper()):
                columnas_fecha.append(columna)
                
            columna += 1
        
        return columnas_fecha

    def convertirXLSB(self,nombre_excelXLSB, extraer_campos = 0):
        wbNew = openpyxl.Workbook()
        hoja = wbNew.active
        with open_xlsb(nombre_excelXLSB) as wb:
            with wb.get_sheet(2) as sheet:
                fila = 0
                columnas_fecha = []
                for row in sheet.rows():
                    row_aux = [None if item.v == None else str(item.v).replace('', ' ').encode(encoding="utf-8",errors="xmlcharrefreplace") for item in row]
                    
                    if(fila==0):
                        campos_tabla = [None if item.v == None else str(item.v).replace('', ' ').encode(encoding="utf-8",errors="xmlcharrefreplace").decode("utf-8") for item in row]
                        if(extraer_campos == 1):                    
                            self.crear_tabla(campos_tabla)
                            exit()
                        else:
                            campos_tabla_insert = campos_tabla
                    if(fila==1):
                        print('Empieza lectura...')
                        x = datetime.datetime.now()
                        print("Tiempo de inicio: "+x.strftime("%X"))

                        columnas_fecha = self.extraerColumnasFechas(row_aux)
                    else:
                        self.conversionFechasFila(columnas_fecha, row_aux)

                    try:
                        hoja.append(row_aux)
                        if(fila>0):
                            row_aux = [None if item.v == None else str(item.v).replace('', ' ').encode(encoding="utf-8",errors="xmlcharrefreplace").decode("utf-8") for item in row]
                            # print(row_aux)
                            self.insertar_registros_tabla(campos_tabla_insert, row_aux)
                    except openpyxl.utils.exceptions.IllegalCharacterError:
                        print('error de escritura para la fila:'+str(fila))
                        quit()
                    
                    fila+=1

        self.guardarArchivoNuevo(wbNew, nombre_excelXLSB)
        print('Termina lectura...')
        x = datetime.datetime.now()
        print("Tiempo de fin: "+x.strftime("%X"))

    def conversionFechasFila(self,columnafecha=[], row=[]):
        if(len(columnafecha) > 0):
            for cf in columnafecha:
                
                if( row[cf] != None and row[cf].decode('utf-8').replace('.', '').isnumeric()):
                    
                    row[cf]=convert_date(float(row[cf].decode('utf-8'))).strftime("%Y-%m-%d").encode(encoding="utf-8") 
                    # print(row[cf])

    def convertirReparto(self,nombre_excelXLSB):
        wbNew = openpyxl.Workbook()
        hoja = wbNew.active
        with open_xlsb(nombre_excelXLSB) as wb:
            with wb.get_sheet(2) as sheet:
                fila = 1
                columnas_fecha = []
                for row in sheet.rows():
                    row_aux = [str(item.v).encode(encoding="utf-8",errors="xmlcharrefreplace") for item in row]
                    row_aux.pop()
                    
                    # if(fila==0):
                        #     campos_tabla = [str(item.v).encode(encoding="utf-8",errors="xmlcharrefreplace").decode("utf-8") for item in row]
                        #     campos_tabla.pop()
                        #     if(extraer_campos == 1):                    
                        #         self.crear_tabla(campos_tabla)
                        #         exit()
                        #     else:
                        #         campos_tabla_insert = campos_tabla
                            
                    if(fila==1):
                        print('Empieza lectura...')
                        x = datetime.datetime.now()
                        print("Tiempo de inicio: "+x.strftime("%X"))

                        columnas_fecha = self.extraerColumnasFechas(row_aux)
                    else:
                        self.conversionFechasFila(columnas_fecha, row_aux)

                    try:
                        hoja.append(row_aux)
                        # row_aux = [str(item.v).encode(encoding="utf-8",errors="xmlcharrefreplace").decode("utf-8") for item in row]
                            # row_aux.pop()
                            # print(row_aux)
                            # self.insertar_registros_tabla(campos_tabla_insert, row_aux)
                    except openpyxl.utils.exceptions.IllegalCharacterError:
                        print('error de escritura para la fila:'+str(fila))
                        quit()
                    
                    fila+=1

        self.guardarArchivoNuevo(wbNew, nombre_excelXLSB)
        print('Termina lectura...')
        x = datetime.datetime.now()
        print("Tiempo de fin: "+x.strftime("%X"))

    def guardarArchivoNuevo(self,wbNew, nombre=""):
        nombre_guardar = path.basename(nombre).replace('.xlsb', '.xlsx')
        
        try:
            if(sys.argv[3]):
                nombre_guardar = sys.argv[3]
        except IndexError:
            pass

        try:
            wbNew.save(nombre_guardar)            
            print("-----------------------------------------------")
            print("Conversión de xlsb a xlsx exitosa!")
            print("-----------------------------------------------")
        except PermissionError:
            print("")
            print("###############################################################")
            print("Tiene en ejecución un archivo excel con el mismo nombre, cerrar")
            print("Nombre del archivo en ejecución: {}".format(nombre_guardar))
            print("###############################################################")
            print("")
        except FileNotFoundError:
            print("No se pudo acceder al directorio '{}' ".format(path.dirname(nombre_guardar)))
        
    def clear(self):
        _ = system('cls' if name == 'nt' else 'clear')

    def conexion(self):
        try:
            conexion = pymysql.connect(
                                        host='localhost',
                                        user='root',
                                        password='',
                                        db= self.BD
                                    )
            # print("Conexión correcta")
        except (pymysql.err.OperationalError, pymysql.err.InternalError) as e:
            print("Ocurrió un error al conectar: ", e)

        return conexion

    def existe_tabla(self):
        conx = self.conexion()
        with conx.cursor() as cursor:
            sql = "SELECT COUNT(*) AS cantidad FROM information_schema.tables WHERE 1 AND table_schema = '"+str(self.BD)+"' AND table_name = '"+str(self.tabla)+"' ;"
            cursor.execute(sql)
            fila = cursor.fetchone()
            
            if(fila[0] > 0):
                print("La tabla '"+self.tabla+"' se encuentra registrada, se empezará a convertir e insertar los datos del arhivo a la Base de Datos...\n")
                self.convertirXLSB(sys.argv[1])
            else:
                print("\nCreando la tabla '"+self.tabla+"'...")
                self.convertirXLSB(sys.argv[1], 1)
                
    def crear_tabla(self,campos_tabla):
        conx = self.conexion()
        with conx.cursor() as cursor:
            # print(campos_tabla)
            sql = "CREATE TABLE "+self.BD+"."+self.tabla+" (`id` INT(10) UNSIGNED NOT NULL AUTO_INCREMENT ";

            if len(campos_tabla) > 0:        
                for campo in campos_tabla:
                    # print(campo)
                    
                    sql += ", "+str(campo).replace(' ','_')+" VARCHAR(255) DEFAULT NULL "
                        # sql += ", "+campo+" date "
                    
                        
                    # print("\n"+sql)
            sql += ', PRIMARY KEY (`id`) '
            sql += ') ENGINE=MYISAM DEFAULT CHARSET=utf8 '
            cursor.execute(sql)
            conx.commit()
            print("La tabla '"+self.tabla+"' fue creada con exito!")
        self.existe_tabla()

    def insertar_registros_tabla(self,campos_tabla, registros):
        conx = self.conexion()
          
        with conx.cursor() as cursor:
            
            sql = "INSERT INTO "+self.BD+"."+self.tabla+" (id" 
            if (campos_tabla and registros):
                for campos_tabla_aux in campos_tabla:
                    if campos_tabla_aux:
                        sql += ", "+str(campos_tabla_aux).replace(' ', '_')
                    # print(campos_tabla_aux)
                # print("\n"+sql)
                # exit()
            sql += ") "
            sql += "VALUES (NULL"
            
            if (campos_tabla and registros):
                for registros_aux in registros:                
                    if registros_aux:
                        sql += ", '"+str(registros_aux)+"'"
                    else:
                        sql += ", ''"
                        
                
            sql += ")";
            # print("\n"+sql)
            # exit()
            cursor.execute(sql)
            conx.commit()

# INSTANCIA DE LA CLASE CONVERTIDOR:
conver = Convertidor()
 
try:
    if(sys.argv[1] == ''):
        raise IndexError('Nombre del archivo vacio')
        
    funcion='N/D'
    file_name = sys.argv[1]
    try:
        if(sys.argv[2] != '' and sys.argv[2] != '-d'):
            funcion = sys.argv[2]
    except IndexError:
        pass

    if(funcion.upper() == 'REPARTO'):
        ext = path.splitext(sys.argv[1])[1]
        if ext == '.xlsb':
            conver.convertirReparto(sys.argv[1])
            # conver.existe_tabla()
        else:
            print("Error, solo se aceptan archivos .xlsb")

    elif(funcion.upper() == 'N/D'):
        # conver.convertirXLSB(sys.argv[1])
        conver.existe_tabla()
    else:
        print("\n\n")
        print("##############################################")
        print("Advertencia: Opción digitada no se encuentra")
        print("Por favor, especifique la acción a realizar...")
        print("##############################################\n")
        print("-----------------------------------------------")
        print("Opciones:                                     |")
        print("----------------------------------------------|")
        print("REPARTO                                       |")
        print("-----------------------------------------------\n")
except IndexError:
        print('Se requiere el nombre del archivo',IndexError)
